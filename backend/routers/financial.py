import re
import uuid
from datetime import date, timedelta
from fastapi import APIRouter, HTTPException, Request, Depends
from typing import Optional, Dict, Any

try:
    from ..database import get_conn, begin_write, recalculate_bank_balances
    from ..auth import get_current_user, log_audit
    from ..utils import rows_to_dicts, add_months
    from ..schemas.financial import FinancialCreate, FinancialPayRequest, BatchPayRequest, AiLaunchRequest
except ImportError:
    from database import get_conn, begin_write, recalculate_bank_balances
    from auth import get_current_user, log_audit
    from utils import rows_to_dicts, add_months
    from schemas.financial import FinancialCreate, FinancialPayRequest, BatchPayRequest, AiLaunchRequest

router = APIRouter(prefix="/api/financial", tags=["Financeiro"])


@router.get("")
def list_financial_entries(
    type: Optional[str] = None,
    status: Optional[str] = None,
    bank_account_id: Optional[int] = None,
    q: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    current_user: Dict[str, Any] = Depends(get_current_user),
):
    conn = get_conn()
    today_str = date.today().isoformat()
    company_id = current_user.get("company_id", 1)

    sql = """
        SELECT e.*,
               c.name as contact_name, c.whatsapp as contact_whatsapp, c.document as contact_doc,
               cat.code as category_code, cat.name as category_name,
               cc.name as cost_center_name,
               b.name as bank_account_name, b.color as bank_color
        FROM financial_entries e
        LEFT JOIN contacts c ON c.id = e.contact_id
        LEFT JOIN categories cat ON cat.id = e.category_id
        LEFT JOIN cost_centers cc ON cc.id = e.cost_center_id
        LEFT JOIN bank_accounts b ON b.id = e.bank_account_id
        WHERE e.status != 'cancelado' AND e.company_id = ?
    """
    params = [company_id]
    if type and type in ("receita", "despesa"):
        sql += " AND e.type = ?"
        params.append(type)
    if bank_account_id:
        sql += " AND e.bank_account_id = ?"
        params.append(bank_account_id)
    if start_date:
        sql += " AND e.due_date >= ?"
        params.append(start_date)
    if end_date:
        sql += " AND e.due_date <= ?"
        params.append(end_date)
    if q:
        sql += " AND (e.description LIKE ? OR c.name LIKE ? OR e.document_number LIKE ?)"
        params.extend([f"%{q}%", f"%{q}%", f"%{q}%"])

    sql += " ORDER BY e.due_date ASC, e.id DESC"
    rows = rows_to_dicts(conn.execute(sql, params).fetchall())

    filtered = []
    for r in rows:
        net_total = round(float(r["amount"]) + float(r.get("interest") or 0.0) + float(r.get("fine") or 0.0) - float(r.get("discount") or 0.0), 2)
        paid = round(float(r.get("paid_amount") or 0.0), 2)
        remaining = max(0.0, round(net_total - paid, 2))
        r["net_total"] = net_total
        r["remaining_amount"] = remaining

        if r["status"] == "pago" or remaining <= 0.009:
            display_status = "pago"
        elif paid > 0 and remaining > 0:
            display_status = "parcial"
        elif r["due_date"] < today_str:
            display_status = "vencido"
        elif r["due_date"] == today_str:
            display_status = "vence_hoje"
        else:
            display_status = "a_vencer"
        r["display_status"] = display_status

        if status and status != "todos":
            if status == "abertos" and display_status == "pago":
                continue
            elif status != "abertos" and display_status != status:
                continue

        filtered.append(r)

    conn.close()
    return filtered


@router.post("")
def create_financial_entry(req: FinancialCreate, current_user: Dict[str, Any] = Depends(get_current_user)):
    conn = get_conn()
    cur = conn.cursor()
    company_id = current_user.get("company_id", 1)

    e_type = req.type
    description = req.description.strip()
    contact_id = req.contact_id

    if req.category_id:
        category_id = req.category_id
    else:
        cat = cur.execute("SELECT id FROM categories WHERE type = ? ORDER BY code ASC LIMIT 1", (e_type,)).fetchone()
        category_id = cat["id"] if cat else (1 if e_type == "receita" else 11)

    if req.cost_center_id:
        cost_center_id = req.cost_center_id
    else:
        cc = cur.execute("SELECT id FROM cost_centers WHERE active = 1 ORDER BY id ASC LIMIT 1").fetchone()
        cost_center_id = cc["id"] if cc else 1

    if req.bank_account_id:
        bank_account_id = req.bank_account_id
    else:
        ba = cur.execute("SELECT id FROM bank_accounts WHERE active = 1 AND company_id = ? ORDER BY id ASC LIMIT 1", (company_id,)).fetchone()
        bank_account_id = ba["id"] if ba else 1

    payment_method = req.payment_method or "Pix"
    competence_date = req.competence_date or date.today().isoformat()
    first_due_str = req.due_date or date.today().isoformat()
    first_due = date.fromisoformat(first_due_str[:10])
    total_amount = float(req.amount or 0.0)
    installments = max(1, req.installments or 1)
    recurrence = req.recurrence or "unico"
    already_paid = bool(req.already_paid)
    doc_num = req.document_number or ""
    notes = req.notes or ""

    group_uid = f"FIN-{uuid.uuid4().hex[:8].upper()}"
    per_installment = round(total_amount / installments, 2) if recurrence == "unico" else total_amount

    created_ids = []
    for i in range(1, installments + 1):
        due_d = add_months(first_due, i - 1).isoformat()
        amt = per_installment
        if recurrence == "unico" and i == installments:
            amt = round(total_amount - per_installment * (installments - 1), 2)

        desc_final = f"{description} ({i}/{installments})" if installments > 1 else description
        is_paid_now = already_paid and i == 1
        status = "pago" if is_paid_now else "pendente"
        paid_val = amt if is_paid_now else 0.0

        cur.execute("""
            INSERT INTO financial_entries (
                company_id, type, description, contact_id, category_id, cost_center_id, bank_account_id,
                payment_method, competence_date, due_date, amount, paid_amount, status,
                installment_number, installment_total, group_uid, recurrence, document_number, notes
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            company_id, e_type, desc_final, contact_id, category_id, cost_center_id, bank_account_id,
            payment_method, competence_date, due_d, amt, paid_val, status,
            i, installments, group_uid, recurrence, doc_num, notes,
        ))
        eid = cur.lastrowid
        created_ids.append(eid)

        if is_paid_now:
            cur.execute("""
                INSERT INTO financial_payments (
                    entry_id, bank_account_id, payment_date, principal_paid, total_cash_flow, payment_method, notes
                ) VALUES (?, ?, ?, ?, ?, ?, 'Baixado no ato do lançamento')
            """, (eid, bank_account_id, due_d, amt, amt, payment_method))

    recalculate_bank_balances(conn)
    log_audit(conn, action="create", entity="financial_entries", entity_id=created_ids[0] if created_ids else None, new_data=req.model_dump(), user_id=current_user.get("id"))
    conn.close()
    return {"status": "ok", "created_ids": created_ids}


@router.post("/{entry_id}/pay")
def pay_financial_entry(entry_id: int, req: FinancialPayRequest, current_user: Dict[str, Any] = Depends(get_current_user)):
    conn = get_conn()
    begin_write(conn)
    cur = conn.cursor()
    company_id = current_user.get("company_id", 1)
    entry = cur.execute("SELECT * FROM financial_entries WHERE id = ? AND company_id = ?", (entry_id, company_id)).fetchone()
    if not entry:
        conn.close()
        raise HTTPException(status_code=404, detail="Lançamento não encontrado.")
    if entry["status"] == "pago":
        conn.close()
        raise HTTPException(status_code=400, detail="Este lançamento já está totalmente baixado.")

    bank_account_id = req.bank_account_id
    payment_date = req.payment_date or date.today().isoformat()
    payment_method = req.payment_method or entry["payment_method"] or "Pix"
    interest_paid = float(req.interest or 0.0)
    fine_paid = float(req.fine or 0.0)
    discount_applied = float(req.discount or 0.0)

    new_interest = round(float(entry["interest"] or 0.0) + interest_paid, 2)
    new_fine = round(float(entry["fine"] or 0.0) + fine_paid, 2)
    new_discount = round(float(entry["discount"] or 0.0) + discount_applied, 2)
    net_total = round(float(entry["amount"]) + new_interest + new_fine - new_discount, 2)
    already_paid = float(entry["paid_amount"] or 0.0)

    cash_paid = round(float(req.amount), 2)
    new_paid_total = round(already_paid + cash_paid, 2)
    new_status = "pago" if new_paid_total >= (net_total - 0.01) else "parcial"

    cur.execute("""
        UPDATE financial_entries
        SET interest = ?, fine = ?, discount = ?, paid_amount = ?, status = ?, bank_account_id = ?, payment_method = ?
        WHERE id = ?
    """, (new_interest, new_fine, new_discount, new_paid_total, new_status, bank_account_id, payment_method, entry_id))

    cur.execute("""
        INSERT INTO financial_payments (
            entry_id, bank_account_id, payment_date, principal_paid, interest_paid,
            fine_paid, discount_applied, total_cash_flow, payment_method, notes
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        entry_id, bank_account_id, payment_date, cash_paid, interest_paid,
        fine_paid, discount_applied, cash_paid, payment_method,
        req.notes or f"Baixa {'Total' if new_status == 'pago' else 'Parcial'}",
    ))

    recalculate_bank_balances(conn)
    log_audit(conn, action="pay", entity="financial_entries", entity_id=entry_id, new_data=req.model_dump(), user_id=current_user.get("id"))
    conn.close()

    return {
        "status": "ok",
        "entry_status": new_status,
        "paid_amount": new_paid_total,
        "remaining": max(0.0, round(net_total - new_paid_total, 2)),
    }


@router.post("/{entry_id}/reverse")
def reverse_financial_entry(entry_id: int, current_user: Dict[str, Any] = Depends(get_current_user)):
    conn = get_conn()
    begin_write(conn)
    cur = conn.cursor()
    company_id = current_user.get("company_id", 1)
    if not cur.execute("SELECT 1 FROM financial_entries WHERE id = ? AND company_id = ?", (entry_id, company_id)).fetchone():
        conn.close()
        raise HTTPException(status_code=404, detail="Lançamento não encontrado.")
    cur.execute("DELETE FROM financial_payments WHERE entry_id = ?", (entry_id,))
    cur.execute("""
        UPDATE financial_entries
        SET paid_amount = 0.0, interest = 0.0, fine = 0.0, discount = 0.0, status = 'pendente'
        WHERE id = ? AND company_id = ?
    """, (entry_id, company_id))
    recalculate_bank_balances(conn)
    log_audit(conn, action="reverse", entity="financial_entries", entity_id=entry_id, user_id=current_user.get("id"))
    conn.close()
    return {"status": "ok", "message": "Baixa estornada e saldo bancário recalculado!"}


@router.post("/batch-pay")
def batch_pay_entries(req: BatchPayRequest, current_user: Dict[str, Any] = Depends(get_current_user)):
    conn = get_conn()
    begin_write(conn)
    cur = conn.cursor()
    count = 0
    today_iso = req.payment_date or date.today().isoformat()
    company_id = current_user.get("company_id", 1)

    for eid in req.entry_ids:
        entry = cur.execute("SELECT * FROM financial_entries WHERE id = ? AND company_id = ?", (eid, company_id)).fetchone()
        if not entry or entry["status"] == "pago":
            continue
        entry = dict(entry)
        net = round(float(entry["amount"]) + float(entry.get("interest") or 0.0) + float(entry.get("fine") or 0.0) - float(entry.get("discount") or 0.0), 2)
        rem = max(0.0, round(net - float(entry.get("paid_amount") or 0.0), 2))
        if rem <= 0:
            continue
        b_id = req.bank_account_id
        cur.execute("UPDATE financial_entries SET paid_amount = ?, status = 'pago', bank_account_id = ? WHERE id = ? AND company_id = ?", (net, b_id, eid, company_id))
        cur.execute("""
            INSERT INTO financial_payments (entry_id, bank_account_id, payment_date, principal_paid, total_cash_flow, payment_method, notes)
            VALUES (?, ?, ?, ?, ?, ?, 'Baixa em Lote')
        """, (eid, b_id, today_iso, rem, rem, req.payment_method or "Pix"))
        count += 1

    recalculate_bank_balances(conn)
    log_audit(conn, action="batch_pay", entity="financial_entries", new_data=req.model_dump(), user_id=current_user.get("id"))
    conn.close()
    return {"status": "ok", "count": count}


@router.delete("/{entry_id}")
def delete_financial_entry(entry_id: int, current_user: Dict[str, Any] = Depends(get_current_user)):
    conn = get_conn()
    begin_write(conn)
    company_id = current_user.get("company_id", 1)
    if not conn.execute("SELECT 1 FROM financial_entries WHERE id = ? AND company_id = ?", (entry_id, company_id)).fetchone():
        conn.close()
        raise HTTPException(status_code=404, detail="Lançamento não encontrado.")
    conn.execute("DELETE FROM financial_payments WHERE entry_id = ?", (entry_id,))
    conn.execute("DELETE FROM financial_entries WHERE id = ? AND company_id = ?", (entry_id, company_id))
    recalculate_bank_balances(conn)
    log_audit(conn, action="delete", entity="financial_entries", entity_id=entry_id, user_id=current_user.get("id"))
    conn.close()
    return {"status": "ok"}


@router.post("/ai-launch")
def ai_natural_language_launch(req: AiLaunchRequest, current_user: Dict[str, Any] = Depends(get_current_user)):
    prompt_txt = req.text.strip()
    if not prompt_txt:
        raise HTTPException(status_code=400, detail="Digite um comando para o lançamento inteligente.")

    low = prompt_txt.lower()
    e_type = "receita" if any(w in low for w in ("receita", "receber", "recebi", "venda", "entrada", "honorario", "mensalidade")) else "despesa"

    val_match = re.search(r"(?:r\$\s*)?(\d{1,3}(?:\.\d{3})*(?:,\d{1,2})|\d+(?:[.,]\d{1,2})?)", low)
    amount = 100.0
    if val_match:
        raw_v = val_match.group(1)
        if "," in raw_v:
            raw_v = raw_v.replace(".", "").replace(",", ".")
        amount = float(raw_v)

    inst_match = re.search(r"(\d{1,2})\s*(?:x|vezes|parcelas)", low)
    installments = int(inst_match.group(1)) if inst_match else 1

    due_d = date.today()
    if "amanha" in low or "amanhã" in low:
        due_d = date.today() + timedelta(days=1)
    elif dm := re.search(r"dia\s+(\d{1,2})", low):
        day_num = min(28, max(1, int(dm.group(1))))
        due_d = date(date.today().year, date.today().month, day_num)
    elif d_slash := re.search(r"(\d{1,2})/(\d{1,2})", low):
        due_d = date(date.today().year, int(d_slash.group(2)), int(d_slash.group(1)))

    already_paid = any(w in low for w in ("pago", "paguei", "recebido", "recebi", "baixado", "à vista"))

    method = "Pix"
    if "boleto" in low:
        method = "Boleto"
    elif "cartao" in low or "cartão" in low or "credito" in low:
        method = "Cartão de Crédito"
    elif "dinheiro" in low:
        method = "Dinheiro"

    company_id = current_user.get("company_id", 1)
    conn = get_conn()
    cur = conn.cursor()
    categories = rows_to_dicts(cur.execute("SELECT * FROM categories WHERE type = ?", (e_type,)).fetchall())
    cat_id = categories[0]["id"] if categories else 1
    cat_name = categories[0]["name"] if categories else "Geral"
    for c in categories:
        kws = [k.strip().lower() for k in (c.get("keywords") or "").split(",") if k.strip()]
        if any(k in low for k in kws):
            cat_id = c["id"]
            cat_name = c["name"]
            break

    banks = rows_to_dicts(cur.execute("SELECT * FROM bank_accounts WHERE active = 1 AND company_id = ?", (company_id,)).fetchall())
    bank_id = banks[0]["id"] if banks else 1
    for b in banks:
        if b["name"].lower() in low or b["bank_name"].lower() in low:
            bank_id = b["id"]
            break

    desc_clean = prompt_txt[0].upper() + prompt_txt[1:]
    group_uid = f"IA-{uuid.uuid4().hex[:6].upper()}"
    per_inst = round(amount / installments, 2)

    for i in range(1, installments + 1):
        d_iso = add_months(due_d, i - 1).isoformat()
        is_p = already_paid and i == 1
        cur.execute("""
            INSERT INTO financial_entries (
                company_id, type, description, category_id, cost_center_id, bank_account_id,
                payment_method, competence_date, due_date, amount, paid_amount,
                status, installment_number, installment_total, group_uid, origin_type
            ) VALUES (?, ?, ?, ?, 1, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'ia')
        """, (
            company_id,
            e_type,
            f"{desc_clean} ({i}/{installments})" if installments > 1 else desc_clean,
            cat_id,
            bank_id,
            method,
            date.today().isoformat(),
            d_iso,
            per_inst,
            per_inst if is_p else 0.0,
            "pago" if is_p else "pendente",
            i,
            installments,
            group_uid,
        ))
        eid = cur.lastrowid
        if is_p:
            cur.execute("""
                INSERT INTO financial_payments (entry_id, bank_account_id, payment_date, principal_paid, total_cash_flow, payment_method, notes)
                VALUES (?, ?, ?, ?, ?, ?, 'Lançamento Inteligente por IA')
            """, (eid, bank_id, d_iso, per_inst, per_inst, method))

    recalculate_bank_balances(conn)
    log_audit(conn, action="ai_launch", entity="financial_entries", new_data={"text": prompt_txt, "amount": amount}, user_id=current_user.get("id"))
    conn.close()

    return {
        "status": "ok",
        "description": desc_clean,
        "amount": amount,
        "type": e_type,
        "category": cat_name,
        "installments": installments,
        "due_date": due_d.isoformat(),
        "already_paid": already_paid,
    }


# ==============================================================================
# IMPORTAÇÃO & EXPORTAÇÃO DE PLANILHAS (EXCEL / CSV)
# ==============================================================================
@router.get("/export-csv")
def export_financial_csv(type: Optional[str] = None, current_user: Dict[str, Any] = Depends(get_current_user)):
    """Exporta lançamentos financeiros em CSV compatível com Excel (ponto e vírgula e UTF-8 BOM)."""
    from fastapi import Response
    import csv
    import io

    conn = get_conn()
    company_id = current_user.get("company_id", 1)
    sql = """
        SELECT e.id, e.type, e.description, c.name as contact, cat.name as category,
               b.name as bank_account, e.due_date, e.competence_date, e.amount,
               e.paid_amount, e.status, e.payment_method
        FROM financial_entries e
        LEFT JOIN contacts c ON c.id = e.contact_id
        LEFT JOIN categories cat ON cat.id = e.category_id
        LEFT JOIN bank_accounts b ON b.id = e.bank_account_id
        WHERE e.status != 'cancelado' AND e.company_id = ?
    """
    params = [company_id]
    if type:
        sql += " AND e.type = ?"
        params.append(type)
    sql += " ORDER BY e.due_date DESC"
    rows = rows_to_dicts(conn.execute(sql, params).fetchall())
    conn.close()

    output = io.StringIO()
    writer = csv.writer(output, delimiter=";", quoting=csv.QUOTE_MINIMAL)
    writer.writerow(["ID", "Tipo", "Descricao", "Contato", "Categoria", "Conta_Bancaria", "Vencimento", "Competencia", "Valor", "Valor_Pago", "Status", "Forma_Pagamento"])

    for r in rows:
        writer.writerow([
            r["id"],
            r["type"].upper(),
            r["description"],
            r.get("contact") or "",
            r.get("category") or "",
            r.get("bank_account") or "",
            r["due_date"],
            r.get("competence_date") or "",
            f"{float(r['amount']):.2f}".replace(".", ","),
            f"{float(r.get('paid_amount') or 0):.2f}".replace(".", ","),
            r["status"].upper(),
            r.get("payment_method") or "",
        ])

    csv_data = "\ufeff" + output.getvalue()
    filename = f"cortex_financeiro_{type or 'geral'}_{date.today().strftime('%Y%m%d')}.csv"
    return Response(
        content=csv_data.encode("utf-8-sig"),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/export-excel")
def export_financial_excel(type: Optional[str] = None, current_user: Dict[str, Any] = Depends(get_current_user)):
    """Exporta lançamentos financeiros em planilha Excel (.xlsx) com openpyxl."""
    from fastapi import Response
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    conn = get_conn()
    company_id = current_user.get("company_id", 1)
    sql = """
        SELECT e.id, e.type, e.description, c.name as contact, cat.name as category,
               b.name as bank_account, e.due_date, e.amount, e.paid_amount, e.status, e.payment_method
        FROM financial_entries e
        LEFT JOIN contacts c ON c.id = e.contact_id
        LEFT JOIN categories cat ON cat.id = e.category_id
        LEFT JOIN bank_accounts b ON b.id = e.bank_account_id
        WHERE e.status != 'cancelado' AND e.company_id = ?
    """
    params = [company_id]
    if type:
        sql += " AND e.type = ?"
        params.append(type)
    sql += " ORDER BY e.due_date DESC"
    rows = rows_to_dicts(conn.execute(sql, params).fetchall())
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Financeiro Cortex"

    headers = ["ID", "Tipo", "Descrição", "Cliente/Fornecedor", "Categoria", "Conta Bancária", "Vencimento", "Valor Total (R$)", "Valor Pago (R$)", "Status", "Forma Pagto"]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")

    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in rows:
        ws.append([
            r["id"],
            r["type"].capitalize(),
            r["description"],
            r.get("contact") or "-",
            r.get("category") or "Geral",
            r.get("bank_account") or "-",
            r["due_date"],
            float(r["amount"]),
            float(r.get("paid_amount") or 0.0),
            r["status"].upper(),
            r.get("payment_method") or "Pix",
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f"cortex_financeiro_{type or 'geral'}_{date.today().strftime('%Y%m%d')}.xlsx"

    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/import-spreadsheet")
async def import_financial_spreadsheet(req: Request, current_user: Dict[str, Any] = Depends(get_current_user)):
    """Importa lançamentos financeiros a partir de um arquivo CSV ou Excel."""
    import csv
    import io
    import openpyxl

    form = await req.form()
    file = form.get("file")
    if not file or not hasattr(file, "read"):
        raise HTTPException(status_code=400, detail="Envie um arquivo CSV ou XLSX válido.")

    filename = getattr(file, "filename", "").lower()
    content = await file.read()
    company_id = current_user.get("company_id", 1)

    imported_count = 0
    records = []

    if filename.endswith(".xlsx") or filename.endswith(".xls"):
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) > 1:
            for r in rows[1:]:
                if not r or len(r) < 3 or not r[1]:
                    continue
                # Espera colunas: [Tipo, Descrição, Valor, Vencimento]
                records.append({
                    "type": str(r[0]).strip().lower() if str(r[0]).strip().lower() in ("receita", "despesa") else "despesa",
                    "description": str(r[1]).strip(),
                    "amount": float(r[2] or 0.0),
                    "due_date": str(r[3])[:10] if r[3] else date.today().isoformat(),
                })
    else:
        # CSV
        text = content.decode("utf-8-sig", errors="ignore")
        delimiter = ";" if ";" in text else ","
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        rows = list(reader)
        if len(rows) > 1:
            for r in rows[1:]:
                if not r or len(r) < 3 or not r[1]:
                    continue
                try:
                    amt_str = str(r[2]).replace(".", "").replace(",", ".") if "," in str(r[2]) else str(r[2])
                    amt = float(amt_str)
                except Exception:
                    amt = 100.0

                records.append({
                    "type": str(r[0]).strip().lower() if str(r[0]).strip().lower() in ("receita", "despesa") else "despesa",
                    "description": str(r[1]).strip(),
                    "amount": amt,
                    "due_date": str(r[3]).strip()[:10] if len(r) > 3 and r[3] else date.today().isoformat(),
                })

    conn = get_conn()
    cur = conn.cursor()
    for rec in records:
        if rec["amount"] <= 0:
            continue
        cur.execute("""
            INSERT INTO financial_entries (company_id, type, description, amount, due_date, competence_date, category_id, cost_center_id, bank_account_id, payment_method, status)
            VALUES (?, ?, ?, ?, ?, ?, 1, 1, 1, 'Boleto', 'pendente')
        """, (company_id, rec["type"], rec["description"], rec["amount"], rec["due_date"], date.today().isoformat()))
        imported_count += 1

    conn.commit()
    log_audit(conn, action="import_spreadsheet", entity="financial_entries", new_data={"imported_count": imported_count}, user_id=current_user.get("id"))
    conn.close()

    return {"status": "ok", "imported_count": imported_count, "message": f"{imported_count} lançamentos importados com sucesso da planilha!"}
